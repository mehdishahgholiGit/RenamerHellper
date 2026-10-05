"""Invoice financial-table extraction (no tkinter). See SPEC.md."""
import os, re
import numpy as np
import cv2
import fitz  # PyMuPDF
import zxingcpp

DPI = 400
PERSIAN = "۰۱۲۳۴۵۶۷۸۹"
ARABIC = "٠١٢٣٤٥٦٧٨٩"
_TR = {c: str(i) for i, c in enumerate(PERSIAN)}
_TR.update({c: str(i) for i, c in enumerate(ARABIC)})


def to_ascii(s):
    return "".join(_TR.get(c, c) for c in s)


def render_page(doc, page_no, dpi=DPI):
    """Render one page to a gray numpy array (PyMuPDF already applies /Rotate)."""
    pix = doc[page_no].get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
    return np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width).copy()


def read_invoice_number(doc):
    """Barcode on page 1 -> digits after the BBFA prefix. Returns (number, raw) or (None, raw)."""
    for dpi in (250, 400):
        pix = doc[0].get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
        arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width)
        for r in zxingcpp.read_barcodes(arr):
            m = re.match(r"^BBFA(\d+)$", r.text.strip())
            if m:
                return m.group(1), r.text
            raw = r.text
            return None, raw
    return None, None


def _group(positions, gap=6):
    out = []
    for p in positions:
        if out and p - out[-1][-1] <= gap:
            out[-1].append(p)
        else:
            out.append([p])
    return [int(np.mean(o)) for o in out]


def find_grid(gray):
    """Return (xs, ys, vmask): vertical/horizontal grid lines of the table.
    Falls back to a lower threshold when a page has only a short (empty) table."""
    H, W = gray.shape
    bw = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                               cv2.THRESH_BINARY_INV, 35, 15)
    h = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (W // 30, 1)))
    hy = h.sum(axis=1) / 255
    ys = _group([i for i in range(H) if hy[i] > W * 0.3])
    xs, v = [], None
    for kdiv, frac in ((40, 0.08), (150, 0.015), (300, 0.008)):
        v = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(10, H // kdiv))))
        vx = v.sum(axis=0) / 255
        xs = _group([i for i in range(W) if vx[i] > H * frac])
        xs = [x for x in xs if 0.04 * W < x < 0.96 * W]
        if len(xs) == 11:
            break
    else:
        xs = []          # caller may reuse the grid of an earlier page
    return xs, ys, v


def _table_rows(xs, ys, vmask):
    """Row bands (y0, y1) whose vertical column lines all span the band."""
    rows = []
    for y0, y1 in zip(ys, ys[1:]):
        if y1 - y0 < 40:
            continue
        ok = 0
        for x in xs:
            band = vmask[y0 + 12:y1 - 12, max(0, x - 4):x + 5]
            if band.size and (band.max(axis=1) > 0).mean() > 0.7:
                ok += 1
        if ok >= len(xs) - 2:
            rows.append((y0, y1))
    return rows


def _ink(cell, inset=10):
    c = cell[inset:-inset, inset:-inset]
    if c.size == 0:
        return 0
    _, bw = cv2.threshold(c, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    return int((bw > 0).sum()) if c.min() < 140 else 0


# column indexes in image order (left -> right)
C_CREDIT, C_DEBIT, C_ROUTE, C_RET, C_FLIGHT, C_ISSUE, C_AIR, C_TICKET, C_NAME, C_ROW = range(10)


def extract_rows(gray, grid):
    """Return (rows, total_row). Each row: dict(cells=[10 gray crops], y=(y0,y1)).
    Data rows have ink in row-number/name/ticket cells; the totals row has none of those
    but ink in debit/credit."""
    xs, ys, vmask = grid
    if len(xs) != 11:
        return [], None
    bands = _table_rows(xs, ys, vmask)[1:]          # first band = header
    cols = list(zip(xs[:-1], xs[1:]))
    data, total = [], None
    for y0, y1 in bands:
        cells = [gray[y0:y1, a:b] for a, b in cols]
        ink = [_ink(c) for c in cells]
        is_data = ink[C_NAME] > 300 and ink[C_TICKET] > 300
        if is_data:
            data.append(dict(cells=cells, y=(y0, y1)))
        elif (ink[C_DEBIT] > 300 or ink[C_CREDIT] > 300) and ink[C_NAME] < 100 and ink[C_TICKET] < 100:
            total = dict(cells=cells, y=(y0, y1))
    return data, total


def find_total_boxes(gray):
    """Crops of the three boxes under the table, as dict debit/credit/balance, or None."""
    H, W = gray.shape
    bw = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                               cv2.THRESH_BINARY_INV, 35, 15)
    cs, _ = cv2.findContours(bw, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    cand = []
    for c in cs:
        x, y, w, h = cv2.boundingRect(c)
        if 0.12 * W < w < 0.16 * W and 0.02 * H < h < 0.03 * H and w / h > 3.5:
            cand.append((x, y, w, h))
    cand.sort(key=lambda t: t[1])
    for i, a in enumerate(cand):
        grp = [b for b in cand if abs(b[1] - a[1]) < 20]
        xs_seen = sorted({b[0] // 40 for b in grp})
        if len(xs_seen) >= 3:
            uniq = []
            for b in sorted(grp, key=lambda t: t[0]):
                if not uniq or b[0] - uniq[-1][0] > 100:
                    uniq.append(b)
                elif b[3] > uniq[-1][3]:
                    uniq[-1] = b
            if len(uniq) == 3:
                crops = [gray[y:y + h, x:x + w] for x, y, w, h in uniq]
                return dict(balance=crops[0], credit=crops[1], debit=crops[2])
    return None


# ------------------------------------------------------------------ digit model
def _app_dir():
    import sys
    if getattr(sys, "frozen", False):
        return sys._MEIPASS
    return os.path.dirname(os.path.abspath(__file__))


def _components(cell, inset):
    """Otsu-binarize a cell and return (bw, [(x, y, w, h)] sorted left->right)."""
    c = cell[inset:-inset, inset:-inset]
    _, bw = cv2.threshold(c, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    n, _, st, _ = cv2.connectedComponentsWithStats(bw, connectivity=8)
    comps = sorted(tuple(int(v) for v in st[i][:4]) for i in range(1, n) if st[i, 4] >= 14)
    return bw, comps


def glyph_feature(bw, box):
    """Scale-invariant glyph descriptor: 12x16 bitmap + aspect ratio."""
    x, y, w, h = box
    g = cv2.resize(bw[y:y + h, x:x + w], (12, 16), interpolation=cv2.INTER_AREA).astype(np.float32) / 255
    return np.concatenate([g.ravel(), [w / h * 6.0]])


class DigitModel:
    """Nearest-neighbour digit classifier for the invoice font (see train_digits.py).
    Families: 'cell' (thin font in amount cells) and 'big' (bold font: tickets, totals boxes)."""

    def __init__(self, path=None):
        path = path or os.path.join(_app_dir(), "digit_model.npz")
        m = np.load(path)
        self.sets = {"cell": (m["X_cell"], m["Y_cell"]), "big": (m["X_big"], m["Y_big"])}

    def rank(self, family, feat):
        """[(digit_char, distance)] best first, one entry per class."""
        X, Y = self.sets[family]
        d = np.linalg.norm(X - feat, axis=1)
        best = {}
        for dist, y in zip(d, Y):
            y = str(y)
            if y not in best or dist < best[y]:
                best[y] = float(dist)
        return sorted(best.items(), key=lambda t: t[1])


# ------------------------------------------------------------------ amounts
def amount_glyphs(cell, model):
    """Ranked candidates per digit glyph of an amount cell (commas dropped). None if empty."""
    if _ink(cell) < 20:
        return None
    bw, comps = _components(cell, 10)
    digs = [b for b in comps if not (b[2] <= 8 and b[3] <= 14)]      # commas are tiny
    if not digs:
        return None
    return [model.rank("cell", glyph_feature(bw, b)) for b in digs]


def solve_column(glyph_rows, target):
    """Pick one digit per glyph (best or runner-up) so the column sums to `target`.
    glyph_rows: per row, list of ranked candidate lists. Returns list of ints, or None."""
    base_digits = [[int(r[0][0]) for r in row] for row in glyph_rows]

    def val(ds):
        return int("".join(map(str, ds))) if ds else 0
    if sum(val(d) for d in base_digits) == target:
        return [val(d) for d in base_digits]
    unc = []                                  # (margin, row, idx, delta, cost)
    for ri, row in enumerate(glyph_rows):
        n = len(row)
        for gi, ranks in enumerate(row):
            if len(ranks) > 1:
                d1, d2 = ranks[0][1], ranks[1][1]
                delta = (int(ranks[1][0]) - int(ranks[0][0])) * 10 ** (n - 1 - gi)
                unc.append((d2 / (d1 + 1e-6), ri, gi, delta, d2 - d1))
    unc.sort()
    unc = unc[:16]
    need = target - sum(val(d) for d in base_digits)
    sums, costs = np.zeros(1, dtype=np.int64), np.zeros(1)
    for _, _, _, delta, cost in unc:
        sums = np.concatenate([sums, sums + delta])
        costs = np.concatenate([costs, costs + cost])
    hit = np.where(sums == need)[0]
    if len(hit) == 0:
        return None
    mask = int(hit[np.argmin(costs[hit])])
    digits = [list(d) for d in base_digits]
    for k, (_, ri, gi, _, _) in enumerate(unc):
        if mask >> k & 1:
            digits[ri][gi] = int(glyph_rows[ri][gi][1][0])
    return [val(d) for d in digits]


# ------------------------------------------------------------------ totals boxes
def read_box(crop, model):
    """Printed total in one of the three boxes. Digits are found by the d,ddd,ddd comma
    pattern counted from the right. Returns int or None."""
    if _ink(crop, 14) < 20:
        return None
    bw, comps = _components(crop, 14)
    n_items = len(comps)
    if n_items == 1:
        return 0                                      # a lone dot = zero
    ok = [n for n in range(2, 16) if n + (n - 1) // 3 == n_items]
    if not ok:
        return None
    digs = [b for k, b in enumerate(comps[::-1]) if k % 4 != 3][::-1]
    commas = [b for k, b in enumerate(comps[::-1]) if k % 4 == 3]
    ref_h = max(b[3] for b in digs)
    if any(b[3] > 0.62 * ref_h for b in commas):
        return None
    s = "".join(model.rank("big", glyph_feature(bw, b))[0][0] for b in digs)
    return int(s)


# ------------------------------------------------------------------ names / tickets
def _find_tesseract():
    """Locate tesseract.exe: bundled copy first, then a typical Windows install."""
    import sys
    bases = []
    if getattr(sys, "frozen", False):
        bases.append(getattr(sys, "_MEIPASS", os.path.dirname(sys.executable)))
        bases.append(os.path.dirname(sys.executable))
    bases.append(os.path.dirname(os.path.abspath(__file__)))
    candidates = [os.path.join(b, "tesseract", "tesseract.exe") for b in bases]
    candidates += [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]
    for c in candidates:
        if c and os.path.isfile(c):
            return c
    return None


def configure_tesseract():
    """Point pytesseract at a known tesseract.exe. Returns the path or None."""
    try:
        import pytesseract
    except ImportError:
        return None
    path = _find_tesseract()
    if not path:
        return None
    pytesseract.pytesseract.tesseract_cmd = path
    tessdata = os.path.join(os.path.dirname(path), "tessdata")
    if os.path.isdir(tessdata):
        os.environ.setdefault("TESSDATA_PREFIX", tessdata)
    return path


configure_tesseract()


def _ocr_eng(cell):
    import pytesseract
    c = cv2.resize(cell[10:-10, 10:-10], None, fx=2, fy=2, interpolation=cv2.INTER_CUBIC)
    c = cv2.copyMakeBorder(c, 30, 30, 30, 30, cv2.BORDER_CONSTANT, value=255)
    return pytesseract.image_to_string(c, lang="eng", config="--psm 7").strip()


def read_name(cell):
    """Passenger name (Latin). Best-effort OCR: it is NOT covered by the totals check."""
    try:
        t = _ocr_eng(cell).replace("|", "I")
    except Exception:
        return ""
    t = re.sub(r"[^A-Za-z/.\- ]", "", t)
    return re.sub(r"\s+", " ", t).strip()


def _parts_from_cuts(bw, box, cuts):
    x, y, w, h = box
    parts = []
    for a, b in zip(cuts, cuts[1:]):
        sub = bw[y:y + h, x + a:x + b]
        ys_, xs_ = np.where(sub > 0)
        if len(xs_) < 5:
            return None
        parts.append((x + a + int(xs_.min()), y + int(ys_.min()),
                      int(xs_.max() - xs_.min() + 1), int(ys_.max() - ys_.min() + 1)))
    return parts


def _split_blob(bw, box, n, model):
    """Split a merged blob into n glyphs. Cut positions near the equal split are tried and the
    combination whose glyphs best match known digit shapes wins (a digit like 3 has inner
    valleys, so ink minima alone are not reliable)."""
    import itertools
    x, y, w, h = box
    win = max(3, int(0.3 * w / n))
    ranges = []
    for i in range(1, n):
        ideal = int(round(w * i / n))
        ranges.append(range(max(4, ideal - win), min(w - 4, ideal + win + 1)))
    best, best_cost = None, 1e18
    for combo in itertools.product(*ranges):
        if any(b - a < 6 for a, b in zip((0,) + combo, combo + (w,))):
            continue
        parts = _parts_from_cuts(bw, box, (0,) + combo + (w,))
        if not parts:
            continue
        cost = sum(model.rank("big", glyph_feature(bw, p))[0][1] for p in parts)
        if cost < best_cost:
            best, best_cost = parts, cost
    return best or []


def _fix_group(bw, g, e, model, glyph_w=21):
    """Make a group of glyph boxes hold exactly e digits by splitting merged blobs."""
    g = list(g)
    if len(g) >= e or not g:
        return g
    n_i = [max(1, int(round(b[2] / glyph_w))) for b in g]
    if sum(n_i) == e:                                   # width says how many digits each blob has
        out = []
        for b, n in zip(g, n_i):
            out += _split_blob(bw, b, n, model) if n > 1 else [b]
        return out
    while len(g) < e:                                   # fallback: halve the widest blob
        k = max(range(len(g)), key=lambda i: g[i][2])
        parts = _split_blob(bw, g[k], 2, model)
        if len(parts) < 2:
            break
        g[k:k + 1] = parts
    return g


def read_ticket(cell, model):
    """Ticket number dddd-ddd-ddd. Returns (text, reliable)."""
    bw, comps = _components(cell, 10)
    dashes = [b for b in comps if b[2] >= 2.5 * b[3]]
    digs = [b for b in comps if not (b[2] >= 2.5 * b[3])]
    groups, cur = [], []
    for b in sorted(comps):
        if b in dashes:
            groups.append(cur)
            cur = []
        else:
            cur.append(b)
    groups.append(cur)
    expected = [4, 3, 3]
    reliable = len(groups) == 3
    if not reliable:                                   # dashes lost: assume plain 10 glyphs
        groups = [digs[:4], digs[4:7], digs[7:]]
    fixed = [_fix_group(bw, g, e, model) for g, e in zip(groups, expected)]
    if [len(g) for g in fixed] != expected:
        reliable = False
    txt = []
    for g in fixed:
        txt.append("".join(model.rank("big", glyph_feature(bw, b))[0][0] for b in g))
    return "-".join(txt), reliable and bool(TICKET_RE.fullmatch("-".join(txt)))


TICKET_RE = re.compile(r"\d{4}-\d{3}-\d{3}")


# ------------------------------------------------------------------ pipeline
from dataclasses import dataclass, field


@dataclass
class InvoiceResult:
    status: str = "FAILED"            # SUCCESS / WARNING / FAILED
    reason: str = ""
    invoice: str = ""
    rows: list = field(default_factory=list)       # dicts: invoice,row,name,ticket,debit,credit
    totals: dict = field(default_factory=dict)     # printed: debit, credit, balance
    notes: list = field(default_factory=list)


def _fail(res, reason):
    res.status, res.reason = "FAILED", reason
    return res


def process_pdf(path, model=None):
    """Read one invoice PDF. Never renames; see rename_invoice()."""
    res = InvoiceResult()
    try:
        doc = fitz.open(path)
    except Exception:
        return _fail(res, "PDF could not be opened")
    try:
        model = model or DigitModel()
        number, raw = read_invoice_number(doc)
        if raw is None:
            return _fail(res, "Barcode not detected")
        if number is None:
            return _fail(res, "Invoice number not detected")
        res.invoice = number

        rows, boxes, last_xs = [], None, None
        for p in range(len(doc)):
            gray = render_page(doc, p)
            xs, ys, vmask = find_grid(gray)
            if len(xs) != 11 and last_xs:
                xs = last_xs                       # empty continuation page: reuse columns
            if len(xs) == 11:
                last_xs = xs
                page_rows, _ = extract_rows(gray, (xs, ys, vmask))
                for r in page_rows:
                    r["page_gray"] = None
                rows += page_rows
            b = find_total_boxes(gray)
            if b:
                boxes = b
        if not rows or not boxes:
            return _fail(res, "Financial table not detected")

        tot = {k: read_box(boxes[k], model) for k in ("debit", "credit", "balance")}
        res.totals = tot
        if None in tot.values():
            return _fail(res, "OCR failed")
        if tot["debit"] - tot["credit"] != tot["balance"]:
            return _fail(res, "Financial total mismatch")

        cols = {}
        for name, ci in (("debit", C_DEBIT), ("credit", C_CREDIT)):
            gl = [amount_glyphs(r["cells"][ci], model) for r in rows]
            if any(g is None for g in gl):
                return _fail(res, "OCR failed")
            vals = solve_column(gl, tot[name])
            if vals is None:
                return _fail(res, "Financial total mismatch")
            cols[name] = vals

        warn = []
        for i, r in enumerate(rows):
            name = read_name(r["cells"][C_NAME])
            ticket, ok = read_ticket(r["cells"][C_TICKET], model)
            if not name:
                warn.append(f"row {i + 1}: empty name")
            if not ok:
                warn.append(f"row {i + 1}: ticket '{ticket}' looks suspicious")
            res.rows.append(dict(invoice=number, row=i + 1, name=name, ticket=ticket,
                                 debit=cols["debit"][i], credit=cols["credit"][i]))
        res.notes = warn
        res.status, res.reason = ("WARNING", "; ".join(warn)) if warn else ("SUCCESS", "")
        return res
    except Exception as e:                       # never let one bad file stop a batch
        return _fail(res, f"OCR failed ({type(e).__name__}: {e})")
    finally:
        doc.close()


def rename_invoice(path, invoice):
    """Rename path -> <invoice>.pdf in the same folder; never overwrite.
    Returns (new_path or None, message)."""
    folder = os.path.dirname(path)
    target = os.path.join(folder, f"{invoice}.pdf")
    if os.path.abspath(target) == os.path.abspath(path):
        return path, "already named"
    if os.path.exists(target):
        return None, f"Duplicate invoice number: {invoice}.pdf already exists"
    os.rename(path, target)
    return target, "renamed"


# ------------------------------------------------------------------ batch + export
def process_folder(folder, model=None, progress=None, should_stop=None):
    """Process every PDF in `folder`. Renames files on SUCCESS/WARNING only and never overwrites.
    Returns a list of dicts: file, status, reason, invoice, new_name, rows.
    progress(i, total, entry) is called after each file."""
    model = model or DigitModel()
    pdfs = sorted(f for f in os.listdir(folder) if f.lower().endswith(".pdf"))
    seen, out = {}, []
    for i, fname in enumerate(pdfs, 1):
        if should_stop and should_stop():
            break
        path = os.path.join(folder, fname)
        r = process_pdf(path, model)
        entry = dict(file=fname, status=r.status, reason=r.reason, invoice=r.invoice,
                     new_name="", rows=[], totals=r.totals)
        if r.status in ("SUCCESS", "WARNING"):
            if r.invoice in seen:
                entry.update(status="FAILED", reason=f"Duplicate invoice number {r.invoice} (also in {seen[r.invoice]})")
            else:
                new, msg = rename_invoice(path, r.invoice)
                if new is None:
                    entry.update(status="FAILED", reason=msg)
                else:
                    seen[r.invoice] = fname
                    entry.update(new_name=os.path.basename(new), rows=r.rows)
        out.append(entry)
        if progress:
            progress(i, len(pdfs), entry)
    return out


def summarize(entries):
    rows = [row for e in entries for row in e["rows"]]
    debit = sum(r["debit"] for r in rows)
    credit = sum(r["credit"] for r in rows)
    return dict(invoices=sum(1 for e in entries if e["rows"]), rows=len(rows),
                debit=debit, credit=credit, balance=debit - credit,
                total=len(entries), ok=sum(1 for e in entries if e["status"] in ("SUCCESS", "WARNING")),
                failed=sum(1 for e in entries if e["status"] == "FAILED"))


def export_excel(path, entries):
    """Write Details / Summary / Files sheets with openpyxl."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Details"
    ws.append(["Invoice", "Row", "Name", "Ticket Number", "Debit", "Credit"])
    for e in entries:
        for r in e["rows"]:
            ws.append([int(r["invoice"]), r["row"], r["name"], r["ticket"], r["debit"], r["credit"]])
    for c in ws[1]:
        c.font = Font(bold=True)
    for col in ("E", "F"):
        for cell in ws[col][1:]:
            cell.number_format = "#,##0"
    for col, w in zip("ABCDEF", (10, 6, 34, 16, 16, 16)):
        ws.column_dimensions[col].width = w
    sm = summarize(entries)
    ws2 = wb.create_sheet("Summary")
    for k, v in (("Total invoices", sm["invoices"]), ("Total rows", sm["rows"]),
                 ("Total debit", sm["debit"]), ("Total credit", sm["credit"]),
                 ("Total balance", sm["balance"]), ("PDFs processed", sm["total"]),
                 ("Succeeded (incl. warnings)", sm["ok"]), ("Failed", sm["failed"])):
        ws2.append([k, v])
    for row in ws2.iter_rows(min_row=3, max_row=5, min_col=2, max_col=2):
        row[0].number_format = "#,##0"
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 18
    ws3 = wb.create_sheet("Files")
    ws3.append(["Original file", "New file", "Invoice", "Status", "Reason"])
    for e in entries:
        ws3.append([e["file"], e["new_name"], e["invoice"], e["status"], e["reason"]])
    for c in ws3[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDE", (34, 16, 10, 10, 60)):
        ws3.column_dimensions[col].width = w
    wb.save(path)
